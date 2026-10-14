"""账号治理的共享实现（管理端「账号管理」页 与 智能助手 共用）。

单一来源，不做第二套口径：
- **导入**：解析（xlsx / docx / csv / txt）→ 表头别名映射 → 逐行校验（valid / dup / error）→ 落库；
- **删除**：账号定位、保护规则（管理员 / 自己 / 课程 owner）、级联清理个人数据；
- 管理端向导端点（/admin/users/import/*、/admin/users/{id}、/admin/users/batch-delete）
  只做上传 / 响应包装，解析校验与删除逻辑全部在这里；
- 智能助手的工具（agent_st/agent/tools/admin_import.py、admin_delete.py）复用同一套
  `build_records` / `validate_records` / `insert_user` / `delete_block_reason` / `delete_user_cascade`
  —— 密码规则、user_id 生成、重名跳过、保护规则、级联范围等行为与页面完全一致。
"""
from __future__ import annotations

import re

from app.middleware.auth import hash_password
from app.models.agent_st import STQuestionDraft
from app.models.ai import ChatMessage, ChatSession
from app.models.alert import Alert, Message
from app.models.checkin import StudyCheckin
from app.models.course import Course, ResourceProgress, ResourceStudyLog
from app.models.graph import LearningPath
from app.models.practice import AnswerRecord, PracticeSession
from app.models.user import TeacherClass, User
from app.models.user_course import UserCourse

MAX_IMPORT_ROWS = 5000          # 向导单次导入上限
MAX_ASSIST_ROWS = 200           # 智能助手「上传文件 → 录入账号」单次上限（更大量走向导）
MAX_BATCH_DELETE = 200          # 单次删除上限（账号管理页批量删除 与 智能助手删号共用）
IMPORT_FIELD_ORDER = ["username", "name", "password", "className", "studentNo", "dept", "title"]

# 表头别名容错：列名（已归一化：去空白/下划线、转小写）命中任一别名即映射到字段
HEADER_ALIASES = {
    "username": ["用户名", "账号", "登录名", "userid", "user", "用户", "login"],
    "name": ["姓名", "名字", "name", "真实姓名"],
    "password": ["密码", "初始密码", "pwd", "password"],
    "className": ["班级", "班级名", "class", "班"],
    "studentNo": ["学号", "编号", "studentno"],
    "dept": ["部门", "院系", "单位", "dept", "department"],
    "title": ["职称", "职务", "头衔", "title"],
}


def norm_header(s: str) -> str:
    return re.sub(r"[\s_]", "", (s or "").lower())


# ---------------------------------------------------------------------------
# ① 解析：文件 → 原始行列（list[list[str]]）
# ---------------------------------------------------------------------------
def rows_from_xlsx(src) -> list:
    from openpyxl import load_workbook

    wb = load_workbook(src, data_only=True, read_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in r]
        if all(x == "" for x in cells):
            continue
        rows.append(cells)
    wb.close()
    return rows


def rows_from_docx(src) -> list:
    from docx import Document

    doc = Document(src)
    rows = []
    if doc.tables:
        for tbl in doc.tables:
            for tr in tbl.rows:
                cells = [c.text.strip() for c in tr.cells]
                if all(x == "" for x in cells):
                    continue
                rows.append(cells)
        if rows:
            return rows
    # 纯文本：按行，行内按显式分隔符切分（避免按空格误切「计算机 2301 班」）
    lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for line in lines:
        parts = [p.strip() for p in re.split(r"[\t,，;；|｜]+", line) if p.strip()]
        if parts:
            rows.append(parts)
    return rows


def rows_from_text(text: str) -> list:
    """纯文本（csv / txt）→ 行列：按行切分，行内按逗号 / 制表符 / 分号等显式分隔符切。"""
    rows = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        parts = [p.strip() for p in re.split(r"[\t,，;；|｜]+", line) if p.strip()]
        if parts:
            rows.append(parts)
    return rows


def rows_from_upload(filename: str, data: bytes) -> list:
    """按扩展名解析上传字节流 → 行列。解析异常向上抛（调用方决定如何提示）。"""
    name = filename or ""
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    if ext == ".xlsx":
        from io import BytesIO
        return rows_from_xlsx(BytesIO(data))
    if ext == ".docx":
        from io import BytesIO
        return rows_from_docx(BytesIO(data))
    if ext in (".csv", ".txt"):
        text = _decode_text(data)
        return rows_from_text(text)
    raise ValueError(f"不支持解析账号表的文件类型：{ext or '(无后缀)'}")


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gbk", errors="ignore")


# ---------------------------------------------------------------------------
# ② 识别：是不是一张"账号表"？推断角色
# ---------------------------------------------------------------------------
def _alias_norm() -> dict:
    return {f: [norm_header(a) for a in als] for f, als in HEADER_ALIASES.items()}


def _header_hits(first_row: list) -> tuple[dict, int]:
    """首行 → {列下标: 字段}，以及命中别名数。"""
    alias = _alias_norm()
    mapping, matched = {}, 0
    for ci, cell in enumerate(first_row or []):
        n = norm_header(cell)
        for field, als in alias.items():
            if n in als:
                mapping[ci] = field
                matched += 1
                break
    return mapping, matched


def looks_like_account_table(rows: list) -> bool:
    """粗判：像账号表才做结构化解析（避免把普通文本附件误当账号表）。"""
    if not rows:
        return False
    _, matched = _header_hits(rows[0])
    if matched >= 2:
        return True
    # 无表头：至少一行有 ≥2 个非空单元（用户名+姓名 起步）
    return any(sum(1 for c in (r or []) if (c or "").strip()) >= 2 for r in rows[:10])


def guess_role(rows: list) -> str:
    """按表头线索推断角色：有班级/学号 → 学生；有部门/职称 → 教师；否则默认学生。"""
    if not rows:
        return "student"
    mapping, _ = _header_hits(rows[0])
    fields = set(mapping.values())
    if fields & {"className", "studentNo"}:
        return "student"
    if fields & {"dept", "title"}:
        return "teacher"
    return "student"


# ---------------------------------------------------------------------------
# ③ 结构：行列 → 字段记录
# ---------------------------------------------------------------------------
def build_records(rows: list, role: str) -> list:
    """把原始行列解析为字段字典列表；首行若像表头则作为表头，否则按角色固定顺序。"""
    first = rows[0]
    header_map, matched = _header_hits(first)
    if matched >= 2:
        data = rows[1:]

        def to_rec(row):
            rec = {f: "" for f in IMPORT_FIELD_ORDER}
            for ci, field in header_map.items():
                if ci < len(row):
                    rec[field] = row[ci]
            return rec
    else:
        order = ["username", "name", "password", "className", "studentNo"] if role == "student" \
            else ["username", "name", "password", "dept", "title"]
        data = rows

        def to_rec(row):
            rec = {f: "" for f in IMPORT_FIELD_ORDER}
            for i, field in enumerate(order):
                if i < len(row):
                    rec[field] = row[i]
            return rec

    records = []
    for i, row in enumerate(data):
        rec = to_rec(row)
        rec["rowIndex"] = i + 1
        records.append(rec)
    return records


# ---------------------------------------------------------------------------
# ④ 校验与落库
# ---------------------------------------------------------------------------
def validate_records(records: list, role: str, default_password: str, db) -> None:
    """就地给每条记录打 status（valid / dup / error）与 error 原因；回写 password（统一或被文件覆盖）。"""
    existing = set(u[0] for u in db.query(User.username).all()) if db is not None else set()
    seen = {}
    for rec in records:
        uname = (rec.get("username") or "").strip()
        name = (rec.get("name") or "").strip()
        pwd = (default_password or "").strip() or (rec.get("password") or "").strip()
        status = "valid"
        err = None
        if not uname:
            err, status = "用户名不能为空", "error"
        elif not name:
            err, status = "姓名不能为空", "error"
        elif len(pwd) < 6:
            err, status = "密码至少 6 位（或在导入时填写统一初始密码）", "error"
        elif uname in seen:
            err, status = f"文件内重复（第 {seen[uname]} 行已出现）", "dup"
        elif uname in existing:
            err, status = "用户名已存在（库中已有，将跳过）", "dup"
        if status == "valid":
            seen[uname] = rec["rowIndex"]
        rec["username"] = uname
        rec["name"] = name
        rec["password"] = pwd
        rec["status"] = status
        rec["error"] = err


def insert_user(db, role: str, username: str, name: str, password: str,
                className: str, studentNo: str, dept: str, title: str) -> str:
    """复用 create_user 的落库逻辑：生成 user_id、哈希密码、avatar_char=姓名首字、按角色取字段。"""
    if role == "teacher":
        uid = "T" + username[:12]
        while db.query(User).filter(User.user_id == uid).first():
            uid = "T" + username[:8] + str(db.query(User).count())[-4:]
    else:
        uid = "S" + username[:12]
        while db.query(User).filter(User.user_id == uid).first():
            uid = "S" + username[:8] + str(db.query(User).count())[-4:]
    user = User(
        user_id=uid,
        username=username,
        password=hash_password(password),
        name=name,
        role=role,
        avatar_char=name[0],
        student_no=(studentNo or "").strip() if role == "student" else "",
        class_name=(className or "").strip() if role == "student" else "",
        title=(title or "").strip() if role == "teacher" else "",
        dept=(dept or "").strip() if role == "teacher" else "",
    )
    db.add(user)
    return uid


# ---------------------------------------------------------------------------
# ⑤ 删除：定位 / 保护规则 / 级联清理
# 账号管理页（单删 + 批删）与智能助手删号工具共用这一份 ——
# 改保护规则或级联范围时只改这里，两边行为自动一致。
# ---------------------------------------------------------------------------
# 级联范围（人话，删号确认卡片与页面提示共用同一句话）
DELETION_SCOPE = ("删除 users 表账号，并级联清理其个人数据：练习与答题记录、资源进度与学习日志、"
                  "学习路径、打卡、AI 会话与消息、预警与消息、出题草稿、课程成员关系与教师班级关联。"
                  "不删课程、资源文件与题库。")


def find_user(db, ident: str):
    """按 userId 或 username 定位用户（userId 优先）。找不到返回 None。"""
    ident = (ident or "").strip()
    if not ident:
        return None
    user = db.query(User).filter(User.user_id == ident).first()
    if user is None:
        user = db.query(User).filter(User.username == ident).first()
    return user


def delete_block_reason(db, user, actor_user_id: str) -> str:
    """能否删除：能删返回空串，不能删返回**完整的人话原因**。

    这是唯一的保护规则来源 —— NotFound / admin / 自己 / 课程 owner 四种情形。
    `actor_user_id` 传当前操作者（管理员本人），用于挡住「删自己」。
    """
    if user is None:
        return "账号不存在"
    if user.role == "admin":
        return "不能删除管理员账号"
    if actor_user_id and user.user_id == actor_user_id:
        return "不能删除当前登录的管理员账号"
    owned = db.query(Course).filter(Course.owner_id == user.user_id).all()
    if owned:
        names = "、".join(c.name for c in owned[:3])
        return (f"该用户是 {len(owned)} 门课程的创建者（owner），"
                f"请先转交或归档课程后再删号：{names}")
    return ""


def delete_user_cascade(db, user) -> dict:
    """删除用户及其个人关联数据（不动课程内容与资源文件），返回各表清理条数。

    调用方负责 **先** 用 `delete_block_reason` 判定能否删；本函数只做删除，
    顺序为先子表后 users，保证外键不失败（不 commit，由调用方决定）。
    """
    uid = user.user_id
    stats = {}

    # AI 会话消息 → 会话
    sess_ids = [r[0] for r in db.query(ChatSession.session_id).filter(ChatSession.user_id == uid).all()]
    if sess_ids:
        stats["chatMessages"] = db.query(ChatMessage).filter(ChatMessage.session_id.in_(sess_ids)).delete(synchronize_session=False)
        stats["chatSessions"] = db.query(ChatSession).filter(ChatSession.user_id == uid).delete(synchronize_session=False)

    # 练习
    ps_ids = [r[0] for r in db.query(PracticeSession.session_id).filter(PracticeSession.user_id == uid).all()]
    if ps_ids:
        stats["answerRecords"] = db.query(AnswerRecord).filter(
            (AnswerRecord.user_id == uid) | (AnswerRecord.session_id.in_(ps_ids))
        ).delete(synchronize_session=False)
    else:
        stats["answerRecords"] = db.query(AnswerRecord).filter(AnswerRecord.user_id == uid).delete(synchronize_session=False)
    stats["practiceSessions"] = db.query(PracticeSession).filter(PracticeSession.user_id == uid).delete(synchronize_session=False)

    # 学习进度 / 路径 / 打卡
    stats["resourceProgress"] = db.query(ResourceProgress).filter(ResourceProgress.user_id == uid).delete(synchronize_session=False)
    stats["studyLogs"] = db.query(ResourceStudyLog).filter(ResourceStudyLog.user_id == uid).delete(synchronize_session=False)
    stats["learningPaths"] = db.query(LearningPath).filter(LearningPath.user_id == uid).delete(synchronize_session=False)
    stats["checkins"] = db.query(StudyCheckin).filter(StudyCheckin.user_id == uid).delete(synchronize_session=False)

    # 预警 / 消息
    stats["alerts"] = db.query(Alert).filter(Alert.user_id == uid).delete(synchronize_session=False)
    stats["messages"] = db.query(Message).filter(
        (Message.user_id == uid) | (Message.from_user == uid)
    ).delete(synchronize_session=False)

    # AI 出题草稿
    stats["drafts"] = db.query(STQuestionDraft).filter(STQuestionDraft.user_id == uid).delete(synchronize_session=False)

    # 课程成员 / 教师班级关联
    stats["userCourses"] = db.query(UserCourse).filter(UserCourse.user_id == uid).delete(synchronize_session=False)
    stats["teacherClasses"] = db.query(TeacherClass).filter(TeacherClass.teacher_user_id == uid).delete(synchronize_session=False)

    db.delete(user)
    return stats


def resolve_for_delete(db, idents, actor_user_id: str) -> dict:
    """把 userId / username 混合列表解析成 `{deletable, skipped}`（**不写库**）。

    给「助手删号」做前置预览，也用于执行前的重新校验 —— 同一份判据，避免两处漂移。
    """
    seen: set[str] = set()
    deletable: list[dict] = []
    skipped: list[dict] = []
    for raw in (idents or []):
        ident = str(raw or "").strip()
        if not ident or ident in seen:
            continue
        seen.add(ident)
        user = find_user(db, ident)
        reason = delete_block_reason(db, user, actor_user_id)
        if reason:
            skipped.append({"ident": ident, "userId": (user.user_id if user else ""),
                            "username": (user.username if user else ""),
                            "name": (user.name if user else ""), "reason": reason})
        else:
            deletable.append({"userId": user.user_id, "username": user.username,
                              "name": user.name, "role": user.role})
    return {"deletable": deletable, "skipped": skipped}
