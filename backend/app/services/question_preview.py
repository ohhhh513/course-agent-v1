"""题目导入的只读解析、规范化和校验；不提交事务、不创建题目。"""
from __future__ import annotations

import csv
from itertools import islice
from io import BytesIO, StringIO
import json
import re

from app.models.graph import GraphNode
from app.services.blank_questions import is_blank_type, pack_blank


def parse_question_text(text):
    text = str(text or "").strip()
    if not text or len(text) > 80000:
        raise ValueError("题目文本为空或超过 80000 字符")
    if text.startswith(("[", "{")):
        data = json.loads(text)
        return data.get("questions", [data]) if isinstance(data, dict) else data
    blocks = re.split(r"\n(?=\s*(?:第\s*)?\d+\s*[.、)）．])", text.replace("\r\n", "\n"))
    questions = []
    for block in blocks:
        block = re.sub(r"^\s*(?:第\s*)?\d+\s*[.、)）．]\s*", "", block.strip())
        # 兼容一行题干、选项、答案；换行后的题目仍使用相同解析逻辑。
        block = re.sub(r"\s+(?=[A-H][.、)）．]\s*)", "\n", block)
        block = re.sub(r"\s+(?=(?:参考答案|答案|解析|分析)\s*[:：])", "\n", block)
        stem, options, answer, analysis = [], [], "", []
        section = "stem"
        for line in block.splitlines():
            line = line.strip()
            if not line:
                continue
            matched = re.match(r"^(?:参考答案|答案|answer)\s*[:：]\s*(.*)", line, re.I)
            if matched:
                answer, section = matched.group(1).strip(), "answer"
                continue
            matched = re.match(r"^(?:解析|分析|analysis)\s*[:：]\s*(.*)", line, re.I)
            if matched:
                analysis.append(matched.group(1)); section = "analysis"
                continue
            matched = re.match(r"^\(?([A-H])\)?[.、)）．:：]\s*(.*)", line)
            if matched:
                options.append({"key": matched.group(1), "text": matched.group(2)}); section = "options"
            elif section == "options" and options:
                options[-1]["text"] += "\n" + line
            elif section == "analysis":
                analysis.append(line)
            elif section == "stem":
                stem.append(line)
        if stem:
            letter = bool(re.fullmatch(r"[A-Ha-h]+", answer or ""))
            if options:
                qtype = "multi" if letter and len(answer) > 1 else "single"
                answer = answer.upper()
            elif letter:
                qtype = "single"
                answer = answer.upper()
            else:
                qtype = "blank"
            questions.append({"stem": "\n".join(stem), "options": options, "answer": answer,
                              "analysis": "\n".join(analysis), "type": qtype})
    return questions


def file_questions(name, content):
    suffix = name.rsplit(".", 1)[-1].lower()
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("文件超过 10MB")
    if suffix in ("xlsx", "csv"):
        if suffix == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                rows = list(islice(workbook.active.iter_rows(values_only=True), 102))
            finally:
                workbook.close()
        else:
            rows = list(islice(csv.reader(StringIO(content.decode("utf-8-sig"))), 102))
        if not rows:
            raise ValueError("题目表为空")
        aliases = {"题干": "stem", "题目": "stem", "答案": "answer", "解析": "analysis", "题型": "type", "难度": "difficulty", "分值": "score"}
        headers = [aliases.get(str(value or "").strip(), str(value or "").strip()) for value in rows[0]]
        result = []
        for row in rows[1:]:
            if not any(value is not None and str(value).strip() for value in row):
                continue
            record = {key: value for key, value in zip(headers, row) if key and value is not None}
            record["options"] = [{"key": key, "text": str(record.pop(key))} for key in "ABCDEFGH" if key in record]
            result.append(record)
        return result
    if suffix in ("jpg", "jpeg", "png", "gif", "webp"):
        raise ValueError("图片题目需先由支持识图的智能体整理为题干、选项和答案，再打开导入预览")
    if suffix in ("txt", "md", "json"):
        return parse_question_text(content.decode("utf-8-sig"))
    if suffix not in ("docx", "pdf"):
        raise ValueError("题目文件支持 TXT、Markdown、JSON、CSV、XLSX、DOCX 和 PDF")
    from app.agent_st.agent.file_text import extract_file_text
    text = extract_file_text(name, content)
    return parse_question_text(text)


def preview_questions(db, course_id, text="", questions=None):
    questions = questions if questions is not None else parse_question_text(text)
    if not isinstance(questions, list) or not 1 <= len(questions) <= 100:
        raise ValueError("每批需要 1～100 道题目")
    nodes = db.query(GraphNode).filter(GraphNode.course_id == course_id).all()
    chapters = {n.id for n in nodes if n.graph_type == "chapter"}
    kps = {n.id for n in nodes if n.graph_type == "knowledge"}
    normalized, errors = [], []
    for index, raw in enumerate(questions):
        issues = []
        if not isinstance(raw, dict):
            normalized.append({}); errors.append({"index": index, "errors": ["题目必须为对象"]}); continue
        raw_answer = str(raw.get("answer") or "").strip()
        question = {"stem": str(raw.get("stem") or raw.get("question") or "").strip(),
                    "answer": raw_answer, "analysis": str(raw.get("analysis") or ""),
                    "type": str(raw.get("type") or "single"), "status": "pending",
                    "chapterId": str(raw.get("chapterId") or raw.get("chapter_id") or ""),
                    "kpIds": raw.get("kpIds") or raw.get("kp_ids") or ([raw["kp_id"]] if raw.get("kp_id") else [])}
        if question["type"] in ("单选", "单选题"):
            question["type"] = "single"
        if question["type"] in ("多选", "多选题"):
            question["type"] = "multi"
        if is_blank_type(question["type"]):
            question["type"] = "blank"
        if not question["stem"] or len(question["stem"]) > 20000:
            issues.append("题干为空或过长")
        if len(question["analysis"]) > 20000:
            issues.append("解析过长")
        if question["type"] not in ("single", "multi", "blank"):
            issues.append("此导入预览支持单选、多选和填空题")
        for key, default, minimum, maximum in (("difficulty", 3, 1, 5), ("score", 5, 1, 100)):
            try:
                value = int(raw.get(key, default))
                if not minimum <= value <= maximum:
                    raise ValueError()
                question[key] = value
            except (ValueError, TypeError):
                question[key] = raw.get(key); issues.append(f"{key} 超出范围")
        if question["type"] == "blank":
            try:
                packed = pack_blank(question["stem"], raw_answer, raw.get("options"))
                question["options"] = packed["options"]
                question["answer"] = packed["display"]
            except ValueError as exc:
                question["options"] = []
                issues.append(str(exc))
        else:
            question["answer"] = raw_answer.upper()
            options = raw.get("options", [])
            if isinstance(options, dict):
                options = [{"key": key, "text": value} for key, value in options.items()]
            if not isinstance(options, list) or not all(isinstance(option, dict) for option in options):
                options = []; issues.append("选项格式不正确")
            question["options"] = [{"key": str(o.get("key") or "").upper(), "text": str(o.get("text") or "")} for o in options]
            letters = [o["key"] for o in question["options"]]
            if not 2 <= len(options) <= 8 or len(set(letters)) != len(letters) or any(not re.fullmatch("[A-H]", letter) for letter in letters) or any(not o["text"].strip() for o in question["options"]):
                issues.append("选项需为 2～8 个不同字母且内容非空")
            answer = re.sub(r"[,，、\s]", "", question["answer"])
            question["answer"] = answer
            if not answer or any(letter not in letters for letter in answer) or (question["type"] == "single" and len(answer) != 1):
                issues.append("答案不对应选项或不符合题型")
        if question["chapterId"] and question["chapterId"] not in chapters:
            issues.append("章节不属于当前课程")
        if not isinstance(question["kpIds"], list) or any(not isinstance(k, str) or k not in kps for k in question["kpIds"]):
            issues.append("知识点不属于当前课程")
        if raw.get("figure_json") or raw.get("figureJson") or raw.get("hasImage"):
            issues.append("含题图结构的题目请在题目编辑器中核对后保存")
        normalized.append(question)
        if issues:
            errors.append({"index": index, "errors": issues})
    return {"questions": normalized, "errors": errors, "total": len(normalized),
            "valid": len(normalized) - len(errors), "ready": not errors}
