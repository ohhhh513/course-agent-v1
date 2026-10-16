"""
管理员接口: /admin/*
仅 role=admin 可访问；创建 / 删除教师、学生账号。
删除时按外键与业务表级联清理个人数据，并保护课程 owner。
"""
from fastapi import APIRouter, Depends, Body, UploadFile, File, Form, Request, HTTPException
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from typing import Optional, List
import io
import uuid as _uuid

from ..database import get_db
from ..models.user import User
from ..models.ai import ChatSession, ChatMessage
from ..middleware.auth import hash_password
from ..dependencies import require_admin
from ..schemas.common import ok, fail
from ..utils import to_user_dict, fmt_dt, loads
from ..services.user_profiles import profile_metadata
# 账号治理（导入 / 删除）的解析、校验、保护规则与级联清理统一由 services.account_import
# 提供 —— 与智能助手的「上传文件 → 录入系统」「删号」工具共用同一套口径，
# 避免密码规则、user_id 生成、重名跳过、删除保护规则等行为出现两份实现。
from ..services.account_import import (MAX_ASSIST_ROWS,
                                       MAX_BATCH_DELETE,
                                       MAX_IMPORT_ROWS,
                                       build_records as _build_records,
                                       delete_block_reason,
                                       delete_user_cascade,
                                       find_user,
                                       guess_role,
                                       insert_user as _insert_user,
                                       looks_like_account_table,
                                       rows_from_upload,
                                       validate_records as _validate_records)


# ===================== 批量导入（xlsx / docx / csv / txt）=====================
async def _read_import_file_async(file: UploadFile):
    filename = file.filename or "upload"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ("xlsx", "docx", "csv", "txt"):
        return None, f"不支持的文件类型：{ext}。支持 .xlsx / .docx / .csv / .txt（旧版 .doc 请另存为 .docx 或 .xlsx）"
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        return None, "文件过大（上限 10MB）"
    # 直接在内存里解析（BytesIO），避免落临时文件在 Windows 上被 openpyxl/docx 占用句柄导致清理失败
    try:
        rows = rows_from_upload(filename, content)
    except Exception as e:  # 解析异常统一转成用户可读错误
        return None, f"文件解析失败：{e}"
    if not rows:
        return None, "文件内容为空或无法识别有效数据行"
    return rows, None


router = APIRouter(prefix="/api/v1/admin", tags=["管理员"])


class CreateUserReq(BaseModel):
    username: str
    password: str
    role: str                       # teacher | student
    name: str
    className: str = ""             # 学生
    dept: str = ""                  # 教师
    title: str = ""
    studentNo: str = ""


@router.get("/users")
def list_users(
    role: str = "",
    keyword: str = "",
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    q = db.query(User).filter(User.role != "admin")
    if role and role in ("teacher", "student"):
        q = q.filter(User.role == role)
    if keyword:
        like = f"%{keyword}%"
        q = q.filter(User.username.like(like) | User.name.like(like))
    rows = q.order_by(User.user_id.asc()).all()
    profiles = profile_metadata(db, [u.user_id for u in rows])
    items = []
    for u in rows:
        d = to_user_dict(u)
        d.update(profiles[u.user_id])
        d["username"] = u.username
        d["className"] = u.class_name or ""
        d["dept"] = u.dept or ""
        items.append(d)
    return ok({"total": len(items), "list": items})


@router.post("/users")
def create_user(
    body: CreateUserReq,
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """创建教师或学生账号"""
    role = (body.role or "").strip()
    if role not in ("teacher", "student"):
        return fail("角色仅支持 teacher / student", 400)
    username = (body.username or "").strip()
    name = (body.name or "").strip()
    password = body.password or ""
    if not username or not name:
        return fail("用户名与姓名不能为空", 400)
    if len(password) < 6:
        return fail("密码至少 6 位", 400)
    if db.query(User).filter(User.username == username).first():
        return fail("用户名已存在", 400)

    if role == "teacher":
        uid = "T" + username[:12]
        while db.query(User).filter(User.user_id == uid).first():
            uid = "T" + username[:8] + str(db.query(User).count())[-4:]
    else:
        uid = "S" + username[:12]
        while db.query(User).filter(User.user_id == uid).first():
            uid = "S" + username[:8] + str(db.query(User).count())[-4:]

    user = User(
        user_id=uid,
        username=username,
        password=hash_password(password),
        name=name,
        role=role,
        avatar_char=name[0],
        student_no=(body.studentNo or "").strip() if role == "student" else "",
        class_name=(body.className or "").strip() if role == "student" else "",
        title=(body.title or "").strip() if role == "teacher" else "",
        dept=(body.dept or "").strip() if role == "teacher" else "",
    )
    db.add(user)
    db.commit()
    return ok({
        "userId": uid,
        "username": username,
        "role": role,
        "name": name,
    })


class UpdateUserReq(BaseModel):
    username: str = ""
    name: str = ""
    password: str = ""              # 留空 = 不修改密码
    className: str = ""             # 学生
    studentNo: str = ""             # 学生
    dept: str = ""                  # 教师
    title: str = ""                 # 教师


@router.put("/users/{user_id}")
def update_user(
    user_id: str,
    body: UpdateUserReq,
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """编辑教师/学生账号（管理员）。

    - 角色不可改（user_id 前缀与角色绑定）
    - 用户名改动会做唯一性校验；姓名改动同步刷新 avatar_char
    - password 留空表示不修改
    """
    user = db.query(User).filter(User.user_id == user_id).first()
    if not user:
        return fail("账号不存在", 404)
    if user.role == "admin":
        return fail("不能编辑管理员账号", 400)

    username = (body.username or "").strip()
    if username and username != user.username:
        if db.query(User).filter(User.username == username, User.user_id != user_id).first():
            return fail("用户名已存在", 400)
        user.username = username

    name = (body.name or "").strip()
    if name:
        user.name = name
        user.avatar_char = name[0]

    if body.password:
        if len(body.password) < 6:
            return fail("密码至少 6 位", 400)
        user.password = hash_password(body.password)

    if user.role == "student":
        user.class_name = (body.className or "").strip()
        user.student_no = (body.studentNo or "").strip()
    else:
        user.dept = (body.dept or "").strip()
        user.title = (body.title or "").strip()

    db.commit()
    d = to_user_dict(user)
    d["username"] = user.username
    d["className"] = user.class_name or ""
    d["dept"] = user.dept or ""
    return ok(d)


class DeleteUserReq(BaseModel):
    confirmUsername: str = ""
    force: bool = False


@router.delete("/users/{user_id}")
def delete_user(
    user_id: str,
    body: DeleteUserReq = Body(default=DeleteUserReq()),
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """删除教师/学生账号（管理员）。

    保护规则（不能删自己 / 管理员 / 课程 owner）与级联范围统一在
    `services.account_import.delete_block_reason / delete_user_cascade`，
    与「批量删除」「智能助手删号」共用同一份口径。
    """
    user = find_user(db, user_id)
    if user is None:
        return fail("账号不存在", 404)
    if body.confirmUsername and body.confirmUsername.strip() != user.username:
        return fail("确认用户名不一致", 400)
    reason = delete_block_reason(db, user, admin.user_id)
    if reason:
        return fail(reason, 400)

    username, role = user.username, user.role
    try:
        from ..services.chat_attachments import delete_chat_attachments, user_session_attachment_ids

        attachment_ids = user_session_attachment_ids(db, user_id)
        deleted = delete_user_cascade(db, user)
        db.commit()
        delete_chat_attachments(attachment_ids)
    except Exception as e:
        db.rollback()
        return fail(f"删除失败：{e}", 500)

    return ok({"userId": user_id, "username": username, "role": role, "deleted": deleted})


class BatchDeleteReq(BaseModel):
    userIds: List[str] = []


@router.post("/users/batch-delete")
def batch_delete_users(
    body: BatchDeleteReq = Body(default=BatchDeleteReq()),
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """批量删除教师/学生账号（管理员）。

    逐行处理，**不做整体回滚**：成功的照常删除并级联清理；被拒的（管理员账号 / 自己 /
    课程创建者 / 账号不存在）与失败的都进报告 —— 与「批量导入」「智能助手删号」同一语义。
    删除走单删完全相同的服务函数（不删课程、资源与题库）。
    """
    raw = [str(x).strip() for x in (body.userIds or []) if str(x).strip()]
    seen: set[str] = set()
    ids = [x for x in raw if not (x in seen or seen.add(x))]
    if not ids:
        return fail("请先勾选要删除的账号", 400)
    if len(ids) > MAX_BATCH_DELETE:
        return fail(f"单次最多删除 {MAX_BATCH_DELETE} 个账号", 400)

    deleted = refused = failed = 0
    items: list[dict] = []
    for uid in ids:
        user = find_user(db, uid)
        reason = delete_block_reason(db, user, admin.user_id)
        if reason:
            refused += 1
            items.append({"userId": uid, "username": (user.username if user else ""),
                          "status": "refused", "reason": reason})
            continue
        try:
            from ..services.chat_attachments import delete_chat_attachments, user_session_attachment_ids

            attachment_ids = user_session_attachment_ids(db, uid)
            stats = delete_user_cascade(db, user)
            db.commit()
            delete_chat_attachments(attachment_ids)
            deleted += 1
            items.append({"userId": uid, "username": user.username, "role": user.role,
                          "status": "deleted", "deleted": stats})
        except Exception as e:  # noqa: BLE001 —— 单行失败不阻塞其余行
            db.rollback()
            failed += 1
            items.append({"userId": uid, "username": user.username, "status": "failed", "reason": str(e)})
    return ok({"deleted": deleted, "refused": refused, "failed": failed, "items": items})


# ===================== 批量导入端点 =====================
@router.get("/users/import/template")
def import_template(
    role: str = "student",
    format: str = "xlsx",
    admin: User = Depends(require_admin),
):
    """下载批量导入模板（.xlsx / .docx），按角色给出对应列。"""
    if role not in ("teacher", "student"):
        return fail("role 仅支持 teacher / student", 400)
    fmt = (format or "xlsx").lower()
    if fmt not in ("xlsx", "docx"):
        return fail("format 仅支持 xlsx / docx", 400)

    if role == "student":
        headers = ["用户名", "姓名", "密码", "班级", "学号"]
        example = ["s_zhangsan", "张三", "123456", "计算机2301班", "20230101"]
        cn = "学生批量导入模板"
    else:
        headers = ["用户名", "姓名", "密码", "部门", "职称"]
        example = ["t_lisi", "李四", "123456", "计算机学院", "副教授"]
        cn = "教师批量导入模板"
    ascii_name = ("student" if role == "student" else "teacher") + "_import_template." + fmt

    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        ws.append(example)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        from docx import Document
        doc = Document()
        doc.add_heading(cn, level=1)
        doc.add_paragraph("请按表格填写，首行为表头，下方每行一人。用户名 / 姓名 / 密码为必填。")
        table = doc.add_table(rows=1, cols=len(headers))
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
        row = table.add_row().cells
        for i, v in enumerate(example):
            row[i].text = v
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    from urllib.parse import quote
    disp = f"attachment; filename={ascii_name}; filename*=UTF-8''{quote(cn + '.' + fmt)}"
    return StreamingResponse(buf, media_type=media, headers={"Content-Disposition": disp})


@router.post("/users/import/preview")
async def import_preview(
    file: UploadFile = File(...),
    role: str = Form("student"),
    defaultPassword: str = Form(""),
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """解析并校验导入文件，返回预览（不写库、不回传明文密码）。"""
    if role not in ("teacher", "student"):
        return fail("role 仅支持 teacher / student", 400)
    rows, perr = await _read_import_file_async(file)
    if perr:
        return fail(perr, 400)
    records = _build_records(rows, role)
    if not records:
        return fail("文件中没有可解析的数据行", 400)
    if len(records) > MAX_IMPORT_ROWS:
        return fail(f"单次导入最多 {MAX_IMPORT_ROWS} 行", 400)
    _validate_records(records, role, defaultPassword, db)
    total = len(records)
    valid = sum(1 for r in records if r["status"] == "valid")
    dup = sum(1 for r in records if r["status"] == "dup")
    err = sum(1 for r in records if r["status"] == "error")
    out = [
        {"rowIndex": r["rowIndex"], "username": r["username"], "name": r["name"],
         "status": r["status"], "error": r["error"]}
        for r in records
    ]
    return ok({"total": total, "valid": valid, "dup": dup, "error": err,
               "role": role, "rows": out})


@router.post("/users/import/commit")
async def import_commit(
    file: UploadFile = File(...),
    role: str = Form("student"),
    defaultPassword: str = Form(""),
    rowIndexes: str = Form(""),
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """真正落库：重新解析+校验文件（不信任前端），逐行写入，有效行照常落库、失败行进报告。"""
    if role not in ("teacher", "student"):
        return fail("role 仅支持 teacher / student", 400)
    rows, perr = await _read_import_file_async(file)
    if perr:
        return fail(perr, 400)
    records = _build_records(rows, role)
    if not records:
        return fail("文件中没有可解析的数据行", 400)
    if len(records) > MAX_IMPORT_ROWS:
        return fail(f"单次导入最多 {MAX_IMPORT_ROWS} 行", 400)
    _validate_records(records, role, defaultPassword, db)

    want = None
    if (rowIndexes or "").strip():
        try:
            want = set(int(x) for x in rowIndexes.split(",") if x.strip())
        except Exception:
            return fail("rowIndexes 格式错误（应为逗号分隔的行号）", 400)

    created = skipped = failed = 0
    errors = []
    for r in records:
        if want is not None and r["rowIndex"] not in want:
            continue
        if r["status"] != "valid":
            skipped += 1
            continue
        try:
            _insert_user(db, role, r["username"], r["name"], r["password"],
                         r["className"], r["studentNo"], r["dept"], r["title"])
            db.commit()
            created += 1
        except Exception as e:
            db.rollback()
            failed += 1
            errors.append({"rowIndex": r["rowIndex"], "username": r["username"], "reason": str(e)})
    return ok({"created": created, "skipped": skipped, "failed": failed, "errors": errors})


# =============================================================================
# 管理员智能体（平台治理流）：/admin/agent/*
# 与 /ai/* 的答疑共用同一套运行层与 SSE 事件契约，差异只有三点：
#   ① 挂 require_admin（教师 / 学生打不进来）；
#   ② **不吃课程上下文** —— 管理员不在任何课程里也能问（`get_current_course_id`
#      会直接 400「尚未创建或加入课程」），会话以 course_id=NULL 落库，与课程会话互不混合；
#   ③ 任务流固定为 admin_ops（5 个系统级只读查询 + 4 个账号治理工具，
#      其中账号录入 / 删除是写操作、必须经确认卡片），模型拿不到任何教学类工具。
# =============================================================================


# 附件上限统一在 services/attachments（与师生答疑共用一份，改口径只改那里）；
# 账号表结构化解析的行数上限见 services.account_import.MAX_ASSIST_ROWS


@router.post("/agent/chat/stream")
async def admin_agent_chat_stream(request: Request, admin: User = Depends(require_admin)):
    """SSE 流式平台治理答疑 —— 事件名 / 字段与 /ai/chat/stream 完全一致。

    请求体支持两种形态：
      - application/json：{sessionId, question}（无附件时的既有调用方式）；
      - multipart/form-data：question / sessionId / files[]（可选多个附件）。

    附件语义 = 「附到当前问题」：服务端抽取纯文本注入**本轮**模型输入，
    不落库正文、不进知识库；消息的 tool_log 只留附件引用及名称、MIME、大小元数据用于历史预览和恢复。
    """
    from sse_starlette.sse import EventSourceResponse
    from starlette.concurrency import run_in_threadpool

    from .ai import ChatReq, agent_sse_events
    from ..services.attachments import AttachmentError, collect_attachments

    ct = (request.headers.get("content-type") or "").lower()
    attachments: list = []
    account_files: list = []

    if ct.startswith("multipart/form-data"):
        try:
            form = await request.form()
        except Exception as exc:  # noqa: BLE001
            return JSONResponse(status_code=400, content=fail(f"表单解析失败：{exc}"))
        question = str(form.get("question") or "").strip()
        session_id = (str(form.get("sessionId") or "").strip() or None)
        raw_files = [f for f in form.getlist("files") if getattr(f, "filename", None)]
        if not question:
            return JSONResponse(status_code=400, content=fail("请输入问题内容后再发送"))
        # 附件抽取走与师生答疑同一份实现（services/attachments）
        try:
            attachments, raw_uploads = await collect_attachments(raw_files)
        except AttachmentError as exc:
            return JSONResponse(status_code=400, content=fail(str(exc)))

        # 结构化预解析：附件里的学生/教师账号表（供「上传文件 → 助手录入系统」）。
        # 不像账号表 / 解析失败就静默跳过，普通附件问答不受影响；
        # 解析结果只放在服务端 extra（**不进提示词、不落库、不写正文**），
        # 由账号导入工具在模型调用时读取，最终随「待批准动作」参数落 tool_log。
        from ..services.chat_attachments import ChatAttachmentStorageError, persist_chat_uploads
        try:
            persist_chat_uploads(raw_uploads, attachments)
        except ChatAttachmentStorageError:
            return JSONResponse(status_code=500, content=fail("附件保存失败，请重新上传"))

        for name, data in raw_uploads:
            try:
                rows = await run_in_threadpool(rows_from_upload, name, data)
                if not rows or not looks_like_account_table(rows):
                    continue
                role = guess_role(rows)
                records = _build_records(rows, role)
                if not records:
                    continue
                account_files.append({
                    "name": name, "role": role,
                    "totalRows": len(records),
                    "truncated": len(records) > MAX_ASSIST_ROWS,
                    "records": records[:MAX_ASSIST_ROWS],
                })
            except Exception:  # noqa: BLE001 —— 不是账号表就直接跳过
                continue
    else:
        try:
            body = await request.json()
        except Exception:  # noqa: BLE001
            return JSONResponse(status_code=400, content=fail("请求体不是合法 JSON"))
        if not isinstance(body, dict):
            return JSONResponse(status_code=400, content=fail("请求体必须是 JSON 对象"))
        question = str(body.get("question") or "").strip()
        session_id = (str(body.get("sessionId") or "").strip() or None)
        if not question:
            return JSONResponse(status_code=400, content=fail("请输入问题内容"))

    chat_req = ChatReq(sessionId=session_id, question=question)
    extra: dict = {}
    if attachments:
        extra["attachments"] = attachments
    if account_files:
        extra["account_import"] = {"files": account_files}
    return EventSourceResponse(
        agent_sse_events(chat_req, admin, course_id="", flow_id="admin_ops", extra_context=extra or None))


class AdminStopReq(BaseModel):
    sessionId: Optional[str] = None


@router.post("/agent/chat/stop")
def admin_agent_chat_stop(
    body: AdminStopReq = Body(default=AdminStopReq()),
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """停止平台治理会话正在进行的生成（协作式「停止思考」）。

    只置一个内存标记：本轮生成在下一个检查点收尾，并把已经生成的部分作为回答落库。
    会话归属校验：非本人的 / 非平台级会话一律 404。
    """
    from ..agent_st.agent import cancel

    sid = (body.sessionId or "").strip()
    if not sid or sid.lower() == "new":
        return ok({"stopped": False})
    if not _own_admin_session(db, admin, sid):
        return fail("会话不存在", 404)
    cancel.request(sid)
    return ok({"stopped": True, "sessionId": sid})


def _own_admin_session(db: Session, admin: User, session_id: str):
    """取「当前管理员 + 平台级（course_id 为空）」的会话；别人的 / 课程会话一律查不到。"""
    return db.query(ChatSession).filter(
        ChatSession.session_id == session_id,
        ChatSession.user_id == admin.user_id,
        ChatSession.course_id.is_(None),
    ).first()


@router.get("/agent/sessions")
def admin_agent_sessions(db: Session = Depends(get_db), admin: User = Depends(require_admin)):
    """当前管理员的平台治理会话列表（不含任何课程会话）。"""
    rows = db.query(ChatSession).filter(
        ChatSession.user_id == admin.user_id,
        ChatSession.course_id.is_(None),
    ).order_by(ChatSession.updated_at.desc()).all()
    items = [
        {
            "sessionId": s.session_id, "title": s.title,
            "time": fmt_dt(s.updated_at, "%m-%d %H:%M"),
            "rounds": s.rounds, "kp": s.kp_name,
        }
        for s in rows
    ]
    return ok({"total": len(items), "list": items})


@router.get("/agent/recommendation-context")
def admin_agent_recommendation_context(
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """Return aggregate platform signals for this admin's personalized suggestions."""
    from ..agent_st import persistence
    from ..agent_st.agent.llm import chat_available

    return ok({
        "studentAccounts": db.query(User).filter(User.role == "student").count(),
        "teacherAccounts": db.query(User).filter(User.role == "teacher").count(),
        "courses": db.query(persistence.Course).count(),
        "questions": db.query(persistence.Question).count(),
        "resources": db.query(persistence.Resource).count(),
        "knowledgePoints": db.query(persistence.KpDetail).count(),
        "alerts": db.query(persistence.Alert).count(),
        "assistantMode": "live" if chat_available() else "demo",
    })


@router.get("/agent/sessions/{session_id}/messages")
def admin_agent_messages(
    session_id: str,
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """治理会话的消息（字段与 /ai/sessions/{id}/messages 一致，含 toolLog）。"""
    if not _own_admin_session(db, admin, session_id):
        return fail("会话不存在", 404)
    rows = db.query(ChatMessage).filter(
        ChatMessage.session_id == session_id,
    ).order_by(ChatMessage.id.asc()).all()
    items = [
        {
            "role": m.role, "method": m.method, "content": m.content,
            "time": m.time_str, "citations": loads(m.citations) or [],
            "toolLog": loads(getattr(m, "tool_log", None)) or [],
        }
        for m in rows
    ]
    return ok(items)


@router.get("/agent/sessions/{session_id}/attachments/{attachment_id}")
def admin_agent_attachment(
    session_id: str,
    attachment_id: str,
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """Return an attachment referenced by this administrator's platform session."""
    from ..services.chat_attachments import find_session_attachment, resolve_chat_attachment

    if not _own_admin_session(db, admin, session_id):
        return JSONResponse(status_code=404, content=fail("会话不存在", 404))
    attachment = find_session_attachment(db, session_id, attachment_id)
    path = resolve_chat_attachment(attachment_id) if attachment else None
    if not attachment or not path:
        return JSONResponse(status_code=404, content=fail("附件不存在或已清理", 404))
    return FileResponse(
        path,
        media_type=attachment.get("mimeType") or "application/octet-stream",
        filename=attachment.get("name") or "attachment",
        headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"},
    )


@router.delete("/agent/sessions/{session_id}")
def admin_agent_delete_session(
    session_id: str,
    db: Session = Depends(get_db),
    admin: User = Depends(require_admin),
):
    """删除自己的治理会话（含全部消息）。"""
    s = _own_admin_session(db, admin, session_id)
    if not s:
        return fail("会话不存在", 404)
    from ..services.chat_attachments import delete_chat_attachments, session_attachment_ids

    attachment_ids = session_attachment_ids(db, session_id)
    db.query(ChatMessage).filter(ChatMessage.session_id == session_id).delete()
    db.delete(s)
    db.commit()
    delete_chat_attachments(attachment_ids)
    return ok({"sessionId": session_id, "deleted": True})


# ---------------------------------------------------------------------------
# 治理写操作的审批闸门（P2c 复用）：与 /agent/actions/* 同一套语义，
# 但作用域是**平台级**（course_id 为空 → 只在 course_id IS NULL 的会话里找动作），
# 挂 require_admin；教师/学生走 /agent/actions/*（课程绑定的那条）。
# ---------------------------------------------------------------------------
def _load_admin_action(store, user_id: str, action_id: str):
    """取出待批准动作；查不到一律 403（查询范围即权限）。"""
    hit = store.find_action(user_id, "", action_id)
    if not hit:
        raise HTTPException(status_code=403, detail="该动作不存在，或不属于你的平台治理会话")
    return hit


def _run_approved_admin_action(admin: User, action_id: str):
    """确认执行：参数以 tool_log 里落库的那份为准，回放带 approval_replay 标记。"""
    from ..agent_st.agent.context import ToolContext
    from ..agent_st.agent.registry import TOOLS, execute, tool_ok
    from ..agent_st.agent.store import AgentStore

    store = AgentStore(user_id=admin.user_id, course_id="")
    message_id, action = _load_admin_action(store, admin.user_id, action_id)
    status = str(action.get("status") or "pending")

    if status == "executed":
        return ok({"actionId": action_id, "status": "executed", "idempotent": True,
                   "result": action.get("result") or {}})
    if status == "rejected":
        raise HTTPException(status_code=409, detail="该动作已被拒绝，不能再次执行")

    tool = str(action.get("tool") or "")
    spec = TOOLS.get(tool)
    if spec is None:
        raise HTTPException(status_code=400, detail=f"该动作引用的工具已不存在：{tool}")

    ctx = ToolContext(
        store=store,
        flow_id=spec.flows[0] if spec.flows else "admin_ops",
        extra={"role": (getattr(admin, "role", "") or "admin").strip()},
        user_id=admin.user_id,
    )
    ctx.turn["approval_replay"] = True
    result = execute(tool, dict(action.get("args") or {}), ctx)

    if not tool_ok(result):
        reason = ""
        if isinstance(result, dict):
            reason = result.get("error") or "；".join(result.get("errors") or []) or ""
        raise HTTPException(status_code=400, detail=f"执行失败：{reason or '未知原因'}")

    # 回传给前端的结果摘要：录入看 created/skipped，删号看 deleted/refused（都是服务端真实计数）
    summary = {k: result.get(k) for k in ("created", "skipped", "failed", "createdUsers",
                                          "deleted", "refused", "deletedUsers")
               if isinstance(result, dict) and k in result}
    store.set_action_status(message_id, action_id, "executed", result=summary)
    created, removed = summary.get("created"), summary.get("deleted")
    if isinstance(created, int):
        message = f"已执行：成功创建 {created} 个账号"
    elif isinstance(removed, int):
        message = f"已执行：成功删除 {removed} 个账号"
    else:
        message = "已执行：该动作已完成"
    return ok({"actionId": action_id, "status": "executed", "idempotent": False,
               "tool": tool, "result": summary, "message": message})


@router.post("/agent/actions/{action_id}/approve")
def admin_agent_approve(
    action_id: str,
    admin: User = Depends(require_admin),
):
    """确认执行某个待批准动作（平台治理会话；如账号批量录入）。"""
    return _run_approved_admin_action(admin, action_id)


@router.post("/agent/actions/{action_id}/reject")
def admin_agent_reject(
    action_id: str,
    admin: User = Depends(require_admin),
):
    """拒绝某个待批准动作：只改状态，不产生任何副作用。"""
    from ..agent_st.agent.store import AgentStore

    store = AgentStore(user_id=admin.user_id, course_id="")
    message_id, action = _load_admin_action(store, admin.user_id, action_id)
    status = str(action.get("status") or "pending")

    if status == "rejected":
        return ok({"actionId": action_id, "status": "rejected", "idempotent": True})
    if status == "executed":
        raise HTTPException(status_code=409, detail="该动作已执行，不能再拒绝")

    store.set_action_status(message_id, action_id, "rejected")
    return ok({"actionId": action_id, "status": "rejected", "idempotent": False})
