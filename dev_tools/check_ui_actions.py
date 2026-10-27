"""三端界面准备、导入预览、归属与只读验收：只使用临时库副本。"""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import shutil
import sqlite3
import sys
import tempfile

REPO = Path(__file__).resolve().parents[1]
os.environ['DEBUG'] = 'false'


def main():
    real = REPO / 'backend/app/data/course_agent.db'
    original = sha256(real.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix='ui-actions-', ignore_cleanup_errors=True) as temp:
        db_file = Path(temp) / 'copy.db'
        shutil.copy2(real, db_file)
        os.environ['DATABASE_URL'] = 'sqlite:///' + db_file.as_posix()
        os.environ['AGENT_DATA_DIR'] = str(Path(temp) / 'agent')
        sys.path.insert(0, str(REPO / 'backend'))
        from app.agent_st.agent.context import ToolContext
        from app.agent_st.agent.ui_actions import action_catalogue, prepare_action
        from app.agent_st.agent.runtime import _run_tool
        from app.agent_st import persistence as p
        from app.services.question_preview import file_questions, preview_questions
        from app.main import app
        from fastapi.testclient import TestClient

        catalogue = action_catalogue()
        ui = (REPO / 'assets/js/shared/agent-ui.js').read_text(encoding='utf-8')
        for role in ('student', 'teacher', 'admin'):
            html = (REPO / f'{role}.html').read_text(encoding='utf-8')
            sources = html + '\n'.join((REPO / path).read_text(encoding='utf-8')
                                      for src in re.findall(r'<script[^>]+src="([^"]+)"', html)
                                      for path in [src.split('?', 1)[0]]
                                      if path.startswith('assets/js/') and (REPO / path).is_file())
            ids = set(re.findall(r'id=["\']([^"\']+)', sources))
            views = set(re.findall(r'data-view="([^"]+)"', html))
            for spec in catalogue.values():
                if spec['role'] != role:
                    continue
                assert not spec['view'] or spec['view'] in views, spec['id']
                assert not spec['handler'] or f"'{spec['handler']}':" in ui, spec['id']
                assert not spec.get('execute') or f"'{spec['execute']}':" in ui, spec['id']
                selectors = [s['selector'] for s in spec['steps']] + [spec['focus'], spec['file_input']]
                selectors += [f['selector'] for f in spec['fields'].values()]
                for selector in selectors:
                    for element_id in re.findall(r'#([A-Za-z][\w-]*)', selector):
                        assert element_id in ids, (spec['id'], element_id)
        print(f'PASS {len(catalogue)} 个入口：真实页面、控件 ID 与处理器一致')

        def db_state():
            with sqlite3.connect(db_file) as conn:
                return list(conn.iterdump())

        with p.SessionLocal() as db:
            contexts = {}
            for role, username in [('student', 'stu07'), ('teacher', 'wangjg'), ('admin', 'admin')]:
                user = db.query(p.User).filter_by(username=username).one()
                member = db.query(p.UserCourse).filter_by(user_id=user.user_id).first()
                contexts[role] = ToolContext(None, 'admin_ops' if role == 'admin' else 'explain',
                    extra={'role': role, 'courseId': member.course_id if member else ''}, user_id=user.user_id)
            teacher, student, admin = contexts['teacher'], contexts['student'], contexts['admin']
            course = teacher.extra['courseId']
            pupil = db.query(p.User).join(p.UserCourse, p.UserCourse.user_id == p.User.user_id).filter(p.UserCourse.course_id == course, p.User.role == 'student').first()
            chapter = db.query(p.GraphNode).filter_by(course_id=course, graph_type='chapter').first()
            knowledge = db.query(p.GraphNode).filter_by(course_id=course, graph_type='knowledge').first()
            assert knowledge, '演示库缺少知识点，无法验收对话出题入口'

        before = db_state()
        def ok(ctx, target, **kw):
            result = prepare_action(ctx, target, **kw)
            assert result.get('ok'), (target, result)
            assert result['ui_action']['auto'] is False and '确认' in result['note'], '所有入口只准备确认，不授权自动执行'
            assert any(entry.get('kind') == 'ui_action' and entry['plan'] == result['ui_action'] for entry in ctx.tool_log)
            return result['ui_action']
        def denied(ctx, target, **kw):
            result = prepare_action(ctx, target, **kw)
            assert result.get('error'), result
            assert 'ui_action' not in result

        plan = ok(student, 'student.course.join', values={'inviteCode': 'TEST123'}, auto=True)
        assert plan['values']['inviteCode'] == 'TEST123'
        ok(student, 'student.messages')
        ok(student, 'student.practice.mode.order', auto=True)
        ok(student, 'student.practice.mode.wrong', auto=True)
        ok(student, 'student.report.create', auto=True)
        ok(teacher, 'teacher.report.create', auto=True)
        created = ok(teacher, 'teacher.question.create', auto=True, values={
            'kpIds': [knowledge.id], 'count': 6, 'types': ['单选', '填空'], 'difficulty': 3})
        assert created['label'] == '预览题目' and created['values']['types'] == ['single', 'blank']
        ok(student, 'student.practice.chapter', object_id=chapter.id, auto=True)
        # 新策略：生成入口也能先准备确认卡；只准备不能写业务数据（末尾状态快照校验）。
        ok(student, 'student.report.create', auto=False)
        ok(teacher, 'teacher.report.create', auto=False)
        ok(teacher, 'teacher.question.create', auto=False, values={'kpIds': [knowledge.id]})
        denied(teacher, 'teacher.question.create', auto=True)
        denied(student, 'teacher.question.create', auto=True, values={'kpIds': [knowledge.id]})
        denied(student, 'teacher.report.create', auto=True)
        denied(student, 'student.practice.answer', auto=True)
        denied(student, 'student.practice.submit', auto=True)
        ok(teacher, 'teacher.monitor.message', object_id=pupil.user_id, values={'content': '请复习树的遍历。'}, auto=True)
        ok(teacher, 'teacher.structure.chapter.edit', object_id=chapter.id, values={'name': '绪论'})
        ok(admin, 'admin.users.new', values={'role': 'teacher', 'name': '测试教师'})
        ok(admin, 'admin.users.import', values={'role': 'student', 'text': '用户名,姓名\ntrial_stu,测试学生'})
        denied(student, 'teacher.monitor.message', object_id=pupil.user_id)
        denied(admin, 'teacher.question.import')
        denied(teacher, 'teacher.monitor.message', object_id=admin.user_id)
        denied(teacher, 'teacher.structure.kp.edit', keyword='不存在的知识点')
        denied(student, 'student.course.join', values={'selector': '#jcGo'})
        denied(student, 'student.course.join', values=[])
        denied(student, 'student.course.join', auto='true')
        denied(teacher, 'teacher.question.new', values={'chapterId': 'OTHER-COURSE'})
        denied(teacher, 'teacher.question.new', values={'options': [{'key': '', 'text': 'a'}, {'key': 'B', 'text': 'b'}]})
        denied(teacher, 'teacher.question.new', use_attachments=True)
        events, result = _run_tool('prepare_ui_action', {'target': 'student.messages'}, student)
        assert [event['type'] for event in events] == ['tool_start', 'tool_end']
        assert set(events[1]) == {'type', 'name', 'ok', 'preview'}
        assert events[1]['preview']['ui_action'] == result['ui_action']
        print('PASS 角色、对象归属、字段白名单、严格布尔值、日志与原有 SSE 字段')

        q = {'stem': '1+1 等于多少？', 'options': [{'key': 'A', 'text': '2'}, {'key': 'B', 'text': '3'}], 'answer': 'A'}
        text = '1. 1+1 等于多少？\nA. 2\nB. 3\n答案：A\n解析：一加一等于二。'
        with p.SessionLocal() as db:
            for name, content in [('example.txt', text.encode()), ('example.md', text.encode()),
                                   ('example.json', json.dumps([q]).encode()),
                                   ('example.csv', '题干,A,B,答案\n1+1,2,3,A'.encode())]:
                preview = preview_questions(db, course, questions=file_questions(name, content))
                assert preview['ready'] and preview['total'] == 1, preview
            from openpyxl import Workbook
            workbook = Workbook(); workbook.active.append(['题干', 'A', 'B', '答案'])
            workbook.active.append(['1+1', '2', '3', 'A'])
            stream = BytesIO(); workbook.save(stream)
            assert preview_questions(db, course, questions=file_questions('example.xlsx', stream.getvalue()))['ready']
            from docx import Document
            doc = Document(); doc.add_paragraph(text)
            stream = BytesIO(); doc.save(stream)
            assert preview_questions(db, course, questions=file_questions('example.docx', stream.getvalue()))['ready']
            from pypdf import PdfWriter
            from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
            writer = PdfWriter(); page = writer.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            content = DecodedStreamObject(); content.set_data(b'BT /F1 12 Tf 20 270 Td (1+1?) Tj 0 -20 Td (A. 2) Tj 0 -20 Td (B. 3) Tj 0 -20 Td (answer: A) Tj ET')
            page[NameObject('/Contents')] = writer._add_object(content)
            stream = BytesIO(); writer.write(stream)
            assert preview_questions(db, course, questions=file_questions('example.pdf', stream.getvalue()))['ready']
            assert not preview_questions(db, course, questions=[{**q, 'answer': 'Z'}])['ready']
            assert not preview_questions(db, course, questions=[{**q, 'chapterId': 'OTHER'}])['ready']
            invalid = preview_questions(db, course, questions=[{**q, 'difficulty': 99}])
            assert not invalid['ready'] and not preview_questions(db, course, questions=invalid['questions'])['ready']
        ok(teacher, 'teacher.question.import', values={'questions': [q]})
        denied(teacher, 'teacher.question.import', values={'questions': [{**q, 'difficulty': 99}]})
        teacher.extra['attachments'] = [{'attachment_id': 'a' * 32, 'name': 'example.txt', 'mime_type': 'text/plain'}]
        assert ok(teacher, 'teacher.question.import', use_attachments=True)['attachments'][0]['attachmentId'] == 'a' * 32
        print('PASS TXT、Markdown、JSON、CSV、XLSX、DOCX、PDF 只读解析，错误重校验和附件引用')
        assert db_state() == before, '准备与预览不能写业务数据库'

        client = TestClient(app)
        try:
            for username, allowed in [('stu07', False), ('wangjg', True)]:
                login = client.post('/api/v1/auth/login', json={'username': username, 'password': '123456'}).json()
                token = login['data']['token']
                route_before = db_state()  # 普通登录可能更新 last_login，预览不能继续写入。
                response = client.post('/api/v1/question/import/preview', data={'text': text},
                    headers={'Authorization': 'Bearer ' + token, 'X-Course-Id': course})
                if allowed:
                    assert response.status_code == 200 and response.json()['data']['ready'], response.text
                else:
                    assert response.status_code == 403, response.text
                assert db_state() == route_before, '预览路由不能写业务数据库'
            login = client.post('/api/v1/auth/login', json={'username': 'admin', 'password': 'admin'}).json()
            headers = {'Authorization': 'Bearer ' + login['data']['token']}
            route_before = db_state()
            response = client.post('/api/v1/admin/users/import/preview', data={'role': 'student', 'defaultPassword': '123456'},
                files={'file': ('accounts.csv', '用户名,姓名\nui_preview_student,测试学生'.encode(), 'text/csv')}, headers=headers)
            assert response.json()['data']['valid'] == 1, response.text
            assert db_state() == route_before, '账号 CSV 预览不能写入'
        finally:
            client.close()
        from app.database import engine
        engine.dispose()
        print('PASS 实际登录与题目预览路由：教师可用、学生403；库副本完整 SQL 内容不变')
    assert sha256(real.read_bytes()).hexdigest() == original
    print('PASS 真实数据库 SHA256 未改变')


if __name__ == '__main__':
    main()
